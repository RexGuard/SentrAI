"""Build SentrAI.pptx (13 slides, 16:9, dark theme) from pitch/SLIDES.md content.

Run make_assets.py first. Usage: python build.py  ->  SentrAI.pptx next to this file.
Sources for every number are in ../research/SOURCES.md.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
ASSETS = HERE / "assets"
OUT = HERE / "SentrAI.pptx"

BG = RGBColor(0x0F, 0x1B, 0x14)
CARD = RGBColor(0x18, 0x2A, 0x1F)
LINE = RGBColor(0x2C, 0x44, 0x35)
GREEN = RGBColor(0x3F, 0xAE, 0x6A)
TEXT = RGBColor(0xF2, 0xF5, 0xF3)
MUTED = RGBColor(0xA9, 0xBB, 0xAF)
AMBER = RGBColor(0xF2, 0xC1, 0x4E)
RED = RGBColor(0xFF, 0x7A, 0x66)
CRIT = RGBColor(0xE6, 0x8A, 0xAE)
WHITE = RGBColor(0xFA, 0xFA, 0xF7)
FONT = "Segoe UI"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)


# ---------------------------------------------------------------- helpers
def new_slide(prs, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    s.notes_slide.notes_text_frame.text = notes
    return s


def box(slide, x, y, w, h, fill=None, line=None, radius=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius:
        shp.adjustments[0] = 0.06
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, runs, size=20, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, dict) runs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = [(para, {})] if isinstance(para, str) else para
        for t, style in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = style.get("font", FONT)
            f.size = Pt(style.get("size", size))
            f.bold = style.get("bold", bold)
            f.color.rgb = style.get("color", color)
    return tb


def title(slide, t, kicker=None):
    if kicker:
        text(slide, MARGIN, Inches(0.45), Inches(11), Inches(0.4), kicker.upper(), size=13, color=GREEN, bold=True)
    text(slide, MARGIN, Inches(0.8), Inches(11.9), Inches(0.9), t, size=34, bold=True)


def bullets(slide, x, y, w, h, items, size=20, gap=10):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.1
        parts = [(item, {})] if isinstance(item, str) else item
        r = p.add_run()
        r.text = "▪  "
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), GREEN
        for t, style in parts:
            r = p.add_run()
            r.text = t
            r.font.name = style.get("font", FONT)
            r.font.size = Pt(style.get("size", size))
            r.font.bold = style.get("bold", False)
            r.font.color.rgb = style.get("color", TEXT)
    return tb


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), x, y, w, h)


def footer(slide, t):
    text(slide, MARGIN, Inches(6.95), Inches(11.9), Inches(0.35), t, size=11, color=MUTED)


def page_no(slide, n):
    text(slide, Inches(12.2), Inches(6.95), Inches(0.5), Inches(0.35), str(n), size=11, color=MUTED,
         align=PP_ALIGN.RIGHT)


B = {"bold": True}


# ---------------------------------------------------------------- slides
def s01_title(prs):
    s = new_slide(prs, "Ishmail: do not read the slide. Go straight into the hook on slide 2.")
    picture(s, ASSETS / "logo.png", Inches(8.9), Inches(1.2), h=Inches(4.6))
    text(s, MARGIN, Inches(2.0), Inches(8), Inches(1.2), "SentrAI", size=72, bold=True)
    text(s, MARGIN, Inches(3.25), Inches(7.8), Inches(1.2),
         "Accountability-based security for small teams", size=26, color=GREEN, bold=True)
    text(s, MARGIN, Inches(4.6), Inches(7.8), Inches(1.0),
         "A sentry doesn't chase you. It just guards the gate.", size=20, color=MUTED)
    text(s, MARGIN, Inches(6.2), Inches(11), Inches(0.5),
         "Cybersecurity + AI  ·  AI For Impact Hackathon 2026  ·  Erick Sientaro · Ishmail · Hozen",
         size=14, color=MUTED)


def s02_hook(prs, n):
    s = new_slide(prs, "Ishmail, about 15 seconds: \"In December 2020, a Singapore tuition platform's own "
                       "security test found a hole in its website. Nobody fixed it. Two months later, 4,625 "
                       "students' details were for sale on the dark web, and the company only found out when "
                       "the regulator told them.\" The point: the alert existed.")
    title(s, "The breach that was warned about", "A real PDPC case")
    cols = [("Dec 2020", "Pentest finds an SQL injection hole", MUTED),
            ("Unfixed", "The developer never patched it", AMBER),
            ("Feb 2021", "4,625 students' data for sale on the dark web", RED),
            ("S$10,000", "PDPC fine. The company learned of it from the regulator", CRIT)]
    cw, gap, y = Inches(2.8), Inches(0.23), Inches(2.1)
    for i, (big, small, col) in enumerate(cols):
        x = MARGIN + i * (cw + gap)
        box(s, x, y, cw, Inches(2.5), CARD, LINE)
        text(s, x + Inches(0.25), y + Inches(0.3), cw - Inches(0.5), Inches(0.8), big, size=28, bold=True, color=col)
        text(s, x + Inches(0.25), y + Inches(1.05), cw - Inches(0.5), Inches(1.3), small, size=16)
    text(s, MARGIN, Inches(4.9), Inches(11.9), Inches(0.7),
         [[("“The vulnerability was left unfixed until the Incident happened.”", {"size": 22})]], color=TEXT)
    text(s, MARGIN, Inches(5.6), Inches(11.9), Inches(0.6),
         [[("Missing response, not missing alert.", {"bold": True, "color": GREEN, "size": 26})]])
    footer(s, "ChampionTutor Inc. (Private Limited), PDPC decision DP-2103-B7984, published 14 Oct 2021, para 2.")
    page_no(s, n)


def s03_problem(prs, n):
    s = new_slide(prs, "Ishmail: \"They already get alerts. The breach happens because nobody acts on them in time.\" "
                       "PDPA figure: up to S$1M, or 10% of Singapore turnover if it exceeds S$10M, whichever is higher (s48J, since 1 Oct 2022).")
    title(s, "The problem", "Why small teams get breached")
    bullets(s, MARGIN, Inches(2.0), Inches(7.4), Inches(4.5), [
        "Singapore SMEs and private schools hold lots of personal data, with one or two IT staff",
        "No 24/7 security team, so alerts pile up and get ignored (alert fatigue)",
        [("PDPA fines: ", B), ("up to S$1M or 10% of Singapore turnover, whichever is higher", {})],
        "After a breach, no one can prove who knew what, and when",
    ], size=21, gap=16)
    picture(s, ASSETS / "bell.png", Inches(8.4), Inches(2.2), w=Inches(4.3))
    footer(s, "PDPA s48J (PDPC Guide on Active Enforcement, Oct 2022). 10% limb applies where Singapore turnover exceeds S$10M.")
    page_no(s, n)


def s04_one_sentence(prs, n):
    s = new_slide(prs, "Ishmail: stress \"your own tolerance line\" and \"temporary\". Humans stay in charge.")
    title(s, "SentrAI in one sentence", "The idea")
    text(s, MARGIN, Inches(1.8), Inches(11.9), Inches(1.0),
         "Scores your risk from 0 to 100 across web, database and OS, warns the people you choose, "
         "and past your own tolerance line applies a temporary, reversible fix, with a tamper-evident record.",
         size=21, color=MUTED)
    steps = [("Warn", "Alert the on-duty operator", AMBER), ("Remind", "SLA missed: remind, copy the lead", AMBER),
             ("Contain", "Temporary block, 2h expiry", RED), ("Report", "Evidence report + audit chain", GREEN)]
    cw, gap, y = Inches(2.65), Inches(0.43), Inches(3.5)
    for i, (big, small, col) in enumerate(steps):
        x = MARGIN + i * (cw + gap)
        box(s, x, y, cw, Inches(2.2), CARD, col)
        text(s, x, y + Inches(0.35), cw, Inches(0.7), big, size=30, bold=True, color=col, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.2), y + Inches(1.2), cw - Inches(0.4), Inches(0.9), small, size=16, align=PP_ALIGN.CENTER)
        if i < 3:
            text(s, x + cw, y + Inches(0.75), gap, Inches(0.6), "→", size=30, color=MUTED, align=PP_ALIGN.CENTER)
    page_no(s, n)


def s05_cia_na(prs, n):
    s = new_slide(prs, "Hozen: be honest. We prove delivery and acknowledgement, not that someone read it. "
                       "The report says \"Ack: none\", not \"they saw it\".")
    title(s, "CIA + NA", "Two principles")
    cols = [("CIA", "Protect the data", ["Confidentiality", "Integrity", "Availability"], GREEN),
            ("NA", "Prove who did what", ["Non-repudiation: every alert and acknowledgement is hash-chained",
                                           "Authentication: every decision is tied to a named operator"], AMBER)]
    cw = Inches(5.8)
    for i, (big, sub, items, col) in enumerate(cols):
        x = MARGIN + i * (cw + Inches(0.3))
        box(s, x, Inches(1.95), cw, Inches(3.6), CARD, col)
        text(s, x + Inches(0.35), Inches(2.15), cw, Inches(0.8), big, size=40, bold=True, color=col)
        text(s, x + Inches(0.35), Inches(2.95), cw, Inches(0.5), sub, size=17, color=MUTED)
        bullets(s, x + Inches(0.35), Inches(3.55), cw - Inches(0.7), Inches(2.0), items, size=18, gap=8)
    text(s, MARGIN, Inches(5.85), Inches(11.9), Inches(0.6),
         [[("Result: ", {"bold": True, "color": GREEN}), ("“I was never notified” is no longer an excuse.", {})]], size=22)
    page_no(s, n)


def s06_index(prs, n):
    s = new_slide(prs, "Erick: executives read 0 to 100 instantly. The curve means stacking incidents can never push it "
                       "past 100. Severity examples: brute force 30, SQL injection 40, shell spawned 60, bulk data dump 70.")
    title(s, "The 0 to 100 risk index", "How risk is scored")
    mono = {"font": "Consolas", "size": 15, "color": TEXT}
    box(s, MARGIN, Inches(1.95), Inches(6.1), Inches(1.75), CARD, LINE)
    text(s, MARGIN + Inches(0.25), Inches(2.1), Inches(5.8), Inches(2.2), [
        [("points = base × confidence × criticality", mono)],
        [("raw    = Σ points + inaction penalty", mono)],
        [("index  = 100 × (1 − e^(−raw / 60))", mono)],
    ], spacing=1.7)
    bullets(s, MARGIN, Inches(4.0), Inches(6.2), Inches(2.3), [
        [("Inaction penalty: ", B), ("+5 per hour unacknowledged, capped at +30", {})],
        [("Bands: ", B), ("Green 0–29 · ", {"color": GREEN}), ("Amber 30–59 · ", {"color": AMBER}),
         ("Red 60–79 · ", {"color": RED}), ("Critical 80–100", {"color": CRIT})],
        [("Threshold ", B), ("is set by each organization", {})],
    ], size=18, gap=8)
    picture(s, ASSETS / "curve.png", Inches(7.1), Inches(1.95), w=Inches(5.55))
    page_no(s, n)


def s07_escalation(prs, n):
    s = new_slide(prs, "Erick: the slope between hour 0 and 6 is entirely caused by inaction. That is the accountability "
                       "story in one picture. Brute force 30 × 0.94 × 1.5 = 42.3 points; SQL injection ≈ 54. Replace with "
                       "the values from the recorded take if they differ.")
    title(s, "Risk grows when nobody acts", "Escalation over time")
    picture(s, ASSETS / "escalation.png", Inches(5.0), Inches(1.75), h=Inches(5.05))
    bullets(s, MARGIN, Inches(2.0), Inches(4.1), Inches(4.8), [
        "Brute force detected: index jumps into Amber",
        "Alert delivered, no acknowledgement: +5 every hour",
        "SLA (2h) breached: reminder, team lead copied",
        "SQL injection lands on a neglected system: crosses 80, Countersign approves a 2h block",
    ], size=18, gap=14)
    page_no(s, n)


def s08_demo(prs, n):
    s = new_slide(prs, "Erick: say if the attack is a replay and if Jev is in fallback mode. Judges respect honesty more than polish. "
                       "Replace the gauge with a screenshot of the dashboard at the moment the gauge crosses 80.")
    title(s, "Live demo", "What you will see")
    bullets(s, MARGIN, Inches(2.0), Inches(7.2), Inches(4.5), [
        "Lab student portal with synthetic member data, all on localhost",
        "Brute force → classified → operator alert ignored",
        "Inaction penalty climbs → SQL injection → crosses 80",
        "Countersign approves → IP blocked for 2h → evidence report → a human decides",
    ], size=21, gap=16)
    picture(s, ASSETS / "gauge.png", Inches(8.4), Inches(2.0), w=Inches(4.3))
    page_no(s, n)


def s09_report(prs, n):
    s = new_slide(prs, "Ishmail: this is for leadership, not to punish. It shows where the process failed so it can be fixed. "
                       "It is fact-based: delivered, acknowledged or not, and what the system did. Use the real report "
                       "from the Reports page for the screenshot if time allows.")
    title(s, "The Security Evidence Report", "For leadership")
    rows = [("Responsible entity", "On-duty operator and shift", TEXT),
            ("SLA violation", "Overdue by 4h 30m (policy 2h)", AMBER),
            ("Timeline", "Detected → alert delivered → reminder → threshold crossed", TEXT),
            ("Acknowledgement", "Ack: none", RED),
            ("Forced action", "IP blocked, 2h expiry, approved by Countersign", TEXT)]
    y = Inches(1.95)
    box(s, MARGIN, y, Inches(11.9), Inches(3.55), CARD, LINE)
    for i, (k, v, col) in enumerate(rows):
        yy = y + Inches(0.2) + i * Inches(0.65)
        text(s, MARGIN + Inches(0.35), yy, Inches(3.3), Inches(0.5), k, size=18, color=MUTED)
        text(s, MARGIN + Inches(3.7), yy, Inches(8), Inches(0.5), v, size=18, bold=(col != TEXT), color=col)
    chain = ["#41 alert delivered", "#42 Countersign approved", "#43 operator rollback"]
    cw = Inches(3.1)
    for i, c in enumerate(chain):
        x = MARGIN + i * (cw + Inches(0.45))
        box(s, x, Inches(5.8), cw, Inches(0.75), None, GREEN)
        text(s, x, Inches(5.8), cw, Inches(0.75), c, size=15, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 2:
            text(s, x + cw, Inches(5.85), Inches(0.45), Inches(0.6), "→", size=22, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, Inches(11.25), Inches(5.8), Inches(1.9), Inches(0.75), "✔ chain valid", size=15, bold=True,
         color=GREEN, anchor=MSO_ANCHOR.MIDDLE)
    page_no(s, n)


def s10_ethics(prs, n):
    s = new_slide(prs, "Hozen: the attacker is stopped by attribution and prosecution, not retaliation. "
                       "There is no self-defence exception in the Computer Misuse Act.")
    title(s, "Ethics: no hack back", "It holds the line and never crosses it")
    bullets(s, MARGIN, Inches(2.0), Inches(7.3), Inches(4.6), [
        [("SentrAI never attacks back.", B)],
        "Hacking back is a crime in Singapore: Computer Misuse Act s3 (access), s5 (modification), s7 (obstruction)",
        "Attacker IPs are often spoofed or hijacked innocent machines",
        [("Our response: ", B), ("block, tarpit, honeypots and honeytokens, evidence to SingCERT or the police", {})],
    ], size=20, gap=16)
    picture(s, ASSETS / "perimeter.png", Inches(8.5), Inches(1.9), h=Inches(4.4))
    footer(s, "Computer Misuse Act 1993 (2020 Rev. Ed.), ss 3, 5, 7: sso.agc.gov.sg/Act/CMA1993")
    page_no(s, n)


def s11_architecture(prs, n, arch_png):
    s = new_slide(prs, "Erick: in the MVP the agents are Python classes in one core service, called in turn; separate "
                       "services are on the roadmap. No agent has a free shell: only allowlisted playbooks with a TTL. "
                       "AI judges and explains. Deterministic code scores and executes.")
    title(s, "Agent architecture", "How it fits together")
    frame_w = Inches(7.9)
    frame_h = Emu(int(frame_w * 1800 / 2800))
    x, y = Emu(int((W - frame_w) / 2)), Inches(1.7)
    box(s, x - Inches(0.1), y - Inches(0.1), frame_w + Inches(0.2), frame_h + Inches(0.2), WHITE)
    picture(s, arch_png, x, y, w=frame_w)
    page_no(s, n)


def s12_roadmap(prs, n):
    s = new_slide(prs, "Ishmail closes: \"Alerts tell you something is wrong. SentrAI makes sure someone answers, and "
                       "proves it when they don't.\"")
    title(s, "Roadmap and team", "What's next")
    cols = [("Now (MVP)", GREEN, ["Web, DB, OS and server log collectors", "Rules → Jev → fallback, accuracy measured",
                                   "AI planner, two-key approval, off switch", "Firewall blocking, dry-run first",
                                   "Tarpit, honeypots, honeytokens", "Evidence report PDF + hash chain"]),
            ("Next", AMBER, ["Live Jev and Telegram tests", "Windows Event Log agent",
                             "Cloud config scan (S3, security groups)", "Teams, Slack and SMS escalation"]),
            ("Later", CRIT, ["Single-binary collectors", "Agents as separate services",
                             "Pilot with a Singapore private school"])]
    cw = Inches(3.8)
    for i, (h, col, items) in enumerate(cols):
        x = MARGIN + i * (cw + Inches(0.25))
        box(s, x, Inches(1.85), cw, Inches(4.0), CARD, col)
        text(s, x + Inches(0.3), Inches(2.0), cw, Inches(0.5), h, size=22, bold=True, color=col)
        bullets(s, x + Inches(0.3), Inches(2.6), cw - Inches(0.5), Inches(2.5), items, size=13, gap=3)
    team = [("Erick Sientaro", "Developer"), ("Ishmail", "CEO, pitch"), ("Hozen", "Slides and sources")]
    for i, (name, role) in enumerate(team):
        x = MARGIN + i * (cw + Inches(0.25))
        box(s, x, Inches(6.05), cw, Inches(0.7), None, LINE)
        text(s, x, Inches(6.05), cw, Inches(0.7),
             [[(name, {"bold": True, "size": 17}), ("  ·  " + role, {"color": MUTED, "size": 15})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    page_no(s, n)


def s13_refs(prs, n):
    s = new_slide(prs, "Optional. Leave on screen at the end or skip in the video.")
    title(s, "References", "Sources")
    bullets(s, MARGIN, Inches(1.9), Inches(11.9), Inches(4.8), [
        "PDPC, Breach of the Protection Obligation by ChampionTutor Inc. (Private Limited), DP-2103-B7984, 14 Oct 2021",
        "PDPC, Guide on Active Enforcement (revised 1 Oct 2022): financial penalties under PDPA s48J",
        "Computer Misuse Act 1993 (2020 Rev. Ed.), sections 3, 5 and 7, Singapore Statutes Online",
        "TypeSafe System One (Jev) documentation, docs.typesafe.ai",
        "Code and plan: github.com/RexGuard/SentrAI  ·  rexguard.github.io/SentrAI",
    ], size=17, gap=12)
    page_no(s, n)


def build(arch_png: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    s01_title(prs)
    s02_hook(prs, 2)
    s03_problem(prs, 3)
    s04_one_sentence(prs, 4)
    s05_cia_na(prs, 5)
    s06_index(prs, 6)
    s07_escalation(prs, 7)
    s08_demo(prs, 8)
    s09_report(prs, 9)
    s10_ethics(prs, 10)
    s11_architecture(prs, 11, arch_png)
    s12_roadmap(prs, 12)
    s13_refs(prs, 13)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    arch = HERE / "architecture.png"
    print("wrote", build(arch))
